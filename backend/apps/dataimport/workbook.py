"""Reading and writing the Excel files of the data import: a template to download, a filled one to read."""

import datetime as dt
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .columns import Column, ImportKind

MAX_ROWS = 5000
# What a filled template of 5 000 rows unpacks to is a few MB; far beyond this is a zip bomb.
MAX_UNPACKED_BYTES = 100 * 1024 * 1024


class WorkbookError(Exception):
    """The file cannot be read as the template of its kind (not an .xlsx, a missing column...)."""


def template_bytes(kind: ImportKind) -> bytes:
    """The empty template: the titles on the first sheet, an "Aide" sheet that explains each column."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = kind.title
    header_fill = PatternFill("solid", fgColor="1B6E8C")
    for index, column in enumerate(kind.columns, start=1):
        cell = sheet.cell(row=1, column=index, value=_title(column))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = max(16, min(36, len(_title(column)) + 4))
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "A2"
    help_sheet = book.create_sheet("Aide")
    help_sheet.append(["Colonne", "Obligatoire", "Explication", "Exemple"])
    for cell in help_sheet[1]:
        cell.font = Font(bold=True)
    for column in kind.columns:
        help_sheet.append([column.title, "oui" if column.required else "", column.help, column.example])
    for letter, width in zip("ABCD", (30, 12, 70, 24), strict=True):
        help_sheet.column_dimensions[letter].width = width
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _title(column: Column) -> str:
    return f"{column.title} *" if column.required else column.title


def _normal(text: str) -> str:
    plain = unicodedata.normalize("NFKD", text.replace("*", "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", plain).strip().casefold()


@dataclass
class SheetRow:
    number: int
    # Raw cell values by column key (only the columns with something in them).
    cells: dict[str, Any] = field(default_factory=dict)


def _refuse_zip_bomb(content: bytes) -> None:
    """An .xlsx is a zip: 5 MB sent can unpack to gigabytes and fill the server's memory (D-141). The sizes the zip
    announces are checked first; a liar is stopped by zipfile itself, which never unpacks past the announced size."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            unpacked = sum(entry.file_size for entry in archive.infolist())
    except (zipfile.BadZipFile, OSError, ValueError):
        return  # not a zip: load_workbook answers that it is not an Excel file
    if unpacked > MAX_UNPACKED_BYTES:
        raise WorkbookError("Le fichier est trop gros une fois décompressé : coupez-le en plusieurs.")


def read_rows(kind: ImportKind, content: bytes) -> list[SheetRow]:
    """The filled rows of the first sheet, matched to the kind's columns by their titles. Empty rows are skipped."""
    _refuse_zip_bomb(content)
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises many types for a file that is not an .xlsx
        raise WorkbookError("Le fichier n'est pas un classeur Excel (.xlsx).") from error
    sheet = book.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise WorkbookError("Le fichier est vide.")
    by_title = {_normal(column.title): column for column in kind.columns}
    positions: dict[int, Column] = {}
    for index, value in enumerate(header):
        if value is not None and (column := by_title.get(_normal(str(value)))):
            positions[index] = column
    missing = [column.title for column in kind.columns if column.required and column not in positions.values()]
    if missing:
        raise WorkbookError("Colonnes obligatoires introuvables : " + ", ".join(missing) + ".")
    result: list[SheetRow] = []
    for number, values in enumerate(rows, start=2):
        cells = {
            column.key: value
            for index, column in positions.items()
            if index < len(values) and (value := values[index]) is not None and str(value).strip() != ""
        }
        if cells:
            result.append(SheetRow(number, cells))
        if len(result) > MAX_ROWS:
            raise WorkbookError(f"Plus de {MAX_ROWS} lignes : coupez le fichier en plusieurs.")
    return result


def text_value(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def decimal_value(value: Any) -> Decimal:
    """A number typed in Excel (or as text, with a comma), or ValueError."""
    if isinstance(value, bool):
        raise ValueError("nombre attendu")
    if isinstance(value, int | float | Decimal):
        return Decimal(str(value))
    try:
        return Decimal(str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", "."))
    except InvalidOperation:
        raise ValueError("nombre attendu") from None


def integer_value(value: Any) -> int:
    number = decimal_value(value)
    if number != number.to_integral_value():
        raise ValueError("nombre entier attendu")
    return int(number)


def date_value(value: Any) -> dt.date:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value).strip()
    for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(text, pattern).date()  # noqa: DTZ007  (a calendar day, no time zone)
        except ValueError:
            continue
    raise ValueError("date attendue (JJ/MM/AAAA)")
