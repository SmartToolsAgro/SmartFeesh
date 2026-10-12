"""Excel files of the reports (P7-T02, P7-T03, D-113): one sheet, a header row and a totals row.

Amounts and kilos stay exact numbers (Decimal) formatted to three decimals, so the accountant can sum them.
"""

from collections.abc import Sequence
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MONEY_FORMAT = "#,##0.000"


def workbook_bytes(
    *, title: str, headers: Sequence[str], rows: Sequence[Sequence[Any]], totals: Sequence[Any] | None = None
) -> bytes:
    sheet = Workbook().active
    assert sheet is not None
    sheet.title = title[:31]
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row in rows:
        sheet.append(list(row))
    if totals is not None:
        sheet.append(list(totals))
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
    for column in sheet.iter_cols(min_row=1):
        width = max((len(str(cell.value)) for cell in column if cell.value is not None), default=8)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(width + 2, 10), 40)
        for cell in column[1:]:
            if isinstance(cell.value, Decimal):
                cell.number_format = MONEY_FORMAT
    sheet.freeze_panes = "A2"
    out = BytesIO()
    sheet.parent.save(out)
    return out.getvalue()
