# ruff: noqa: F811  (fixtures imported from test_reports are used by name)
"""Financial reports (P7-T04, D-116): the month's result, cash flow, balances and unpaid salaries."""

from decimal import Decimal
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.closing import monthly
from apps.core.models import Role
from apps.core.tests.factories import MembershipFactory
from apps.expenses import selectors as expense_selectors
from apps.expenses import services as expenses
from apps.partners.tests.factories import CustomerFactory
from apps.sales import services as sales_services

from .test_reports import (  # noqa: F401  (fixtures)
    TODAY,
    actor,
    client_for,
    cold_room,
    dorade,
    moment,
    organization,
    owner,
    receive,
    sell,
)

pytestmark = pytest.mark.django_db
D = Decimal


def test_the_month_result_is_the_one_of_the_monthly_closing(owner, organization, actor, cold_room, dorade):
    receive(organization, cold_room, dorade, "50", cost="20.000")
    sell(organization, actor, cold_room, dorade, "10", "30.000")
    month = TODAY.replace(day=1)

    answer = client_for(owner).get(reverse("report-month-result"), {"month": month.isoformat()}).json()

    expected = monthly.month_report(organization, month=month)
    assert D(answer["net_result"]) == expected.net_result
    assert (D(answer["revenue"]), D(answer["cogs"])) == (D("300.000"), D("200.000"))
    assert answer["sales_count"] == 1


def test_cash_flow_counts_money_when_it_moved(owner, organization, actor, cold_room, dorade):
    receive(organization, cold_room, dorade, "50")
    sell(organization, actor, cold_room, dorade, "10", "30.000")  # paid in full: 300 in
    expenses.ensure_default_categories(organization=organization)
    expenses.create_expense(
        organization=organization,
        actor=actor,
        spent_on=TODAY,
        category=expense_selectors.category_by_code(organization, "fuel"),
        amount="45.500",
        approve=True,
    )

    answer = client_for(owner).get(reverse("report-cash-flow"), {"group_by": "month"}).json()

    assert [(r["cash_in"], r["cash_out"], r["net"]) for r in answer["rows"]] == [("300.000", "45.500", "254.500")]
    assert (answer["total_in"], answer["total_out"], answer["net"]) == ("300.000", "45.500", "254.500")


def test_customers_who_owe_come_with_their_ageing(owner, organization, actor, cold_room, dorade):
    receive(organization, cold_room, dorade, "50")
    customer = CustomerFactory.create(organization=organization, name="Restaurant Le Port", credit_limit=D("1000"))
    sales_services.create_sale(
        organization=organization,
        actor=actor,
        customer=customer,
        location=cold_room,
        lines=[{"product": dorade, "quantity": "10", "unit_price": "30.000"}],
        sold_at=moment(TODAY),
        payment={"amount_paid": "100.000", "payment_method": "cash"},
        confirm=True,
    )

    answer = client_for(owner).get(reverse("report-customer-balances")).json()

    assert [(r["name"], r["balance"], r["d0_30"]) for r in answer["rows"]] == [
        ("Restaurant Le Port", "200.000", "200.000")
    ]
    assert answer["total_balance"] == "200.000"


def test_nothing_owed_gives_empty_balance_reports(owner):
    client = client_for(owner)

    assert client.get(reverse("report-supplier-balances")).json()["rows"] == []
    assert client.get(reverse("report-unpaid-salaries")).json() == {"rows": [], "total": "0.000"}


def test_only_the_owner_and_the_accountant_read_financial_reports(owner, organization):
    names = (
        "report-month-result",
        "report-cash-flow",
        "report-customer-balances",
        "report-supplier-balances",
        "report-unpaid-salaries",
    )
    accountant = MembershipFactory.create(organization=organization, role=Role.ACCOUNTANT)
    manager = MembershipFactory.create(organization=organization, role=Role.MANAGER)
    driver = MembershipFactory.create(organization=organization, role=Role.DRIVER)

    for name in names:
        assert client_for(owner).get(reverse(name)).status_code == 200
        assert client_for(accountant).get(reverse(name)).status_code == 200
        assert client_for(manager).get(reverse(name)).status_code == 403
        assert client_for(driver).get(reverse(name)).status_code == 403


def test_the_excel_files_open_and_carry_the_figures(owner, organization, actor, cold_room, dorade):
    receive(organization, cold_room, dorade, "50")
    sell(organization, actor, cold_room, dorade, "10", "30.000")
    client = client_for(owner)

    month = client.get(reverse("report-month-result-export"))
    cash = client.get(reverse("report-cash-flow-export"), {"group_by": "day"})

    assert month.status_code == cash.status_code == 200
    sheet = load_workbook(BytesIO(month.content)).active
    assert sheet is not None
    assert [row[0].value for row in sheet.iter_rows(min_row=2, max_row=3)] == [
        "Chiffre d'affaires HT",
        "Coût des marchandises vendues",
    ]
    cash_sheet = load_workbook(BytesIO(cash.content)).active
    assert cash_sheet is not None
    assert cash_sheet.cell(row=cash_sheet.max_row, column=2).value == D("300.000")
