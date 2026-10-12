# ruff: noqa: F811  (fixtures imported from the closing tests are used by name)
"""Partnership report (P7-T05, D-117): capital and profit share per partner over a period."""

import datetime as dt
from decimal import Decimal
from io import BytesIO

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.closing import services as closing
from apps.closing.tests.conftest import (  # noqa: F401
    actor,
    client_for,
    cold_room,
    finished_lot,
    member,
    organization,
    owner,
    partner,
)
from apps.core.models import Role
from apps.partnerships import selectors as partnerships
from apps.reports import partners

pytestmark = pytest.mark.django_db
D = Decimal
TODAY = timezone.localdate()


def settle(organization, actor, batch):
    closing.close_lot(organization=organization, actor=actor, batch_id=batch.pk, costs_confirmed=True)
    return closing.settle_lot(
        organization=organization,
        actor=actor,
        batch_id=batch.pk,
        payments=[{"amount": "1000.000", "method": "cash"}, {"amount": "286.000", "method": "transfer"}],
        notes="",
    )


def test_an_open_partnership_shows_the_capital_put_in_and_still_outstanding(organization, finished_lot, partner):
    report = partners.partner_report(organization, date_from=TODAY.replace(day=1), date_to=TODAY)

    assert [
        (r.name, r.lots, r.contributed, r.returned, r.capital_outstanding, r.share, r.share_due) for r in report.rows
    ] == [(partner.name, 1, D("1000.000"), D("0.000"), D("1000.000"), D("0.000"), D("0.000"))]
    assert report.totals.capital_outstanding == D("1000.000")


def test_a_settled_lot_shows_capital_returned_and_the_share_paid(organization, actor, finished_lot, partner):
    settle(organization, actor, finished_lot)

    row = partners.partner_report(organization, date_from=TODAY.replace(day=1), date_to=TODAY).rows[0]

    # The same balances as the partner screens say.
    balances = partnerships.partner_balances(organization, partner_id=partner.pk)
    assert (
        (row.capital_outstanding, row.share_due)
        == (balances.capital, balances.profit_share)
        == (D("0.000"), D("0.000"))
    )
    assert (row.contributed, row.returned) == (D("1000.000"), D("1000.000"))
    # 40 % of the 715.000 net, all paid out.
    assert (row.share, row.paid, row.uncovered_loss) == (D("286.000"), D("286.000"), D("0.000"))


def test_the_period_limits_the_flows_but_not_the_balances(organization, finished_lot):
    last_year = TODAY - dt.timedelta(days=400)

    report = partners.partner_report(organization, date_from=last_year, date_to=last_year + dt.timedelta(days=10))

    # Nothing moved back then, but the capital still in the company today is shown.
    [row] = report.rows
    assert (row.contributed, row.returned, row.share) == (D("0.000"), D("0.000"), D("0.000"))
    assert row.capital_outstanding == D("1000.000")


def test_the_partner_report_api_is_for_the_owner_and_the_accountant(organization, actor, owner, finished_lot):
    settle(organization, actor, finished_lot)
    accountant = member(organization, Role.ACCOUNTANT)
    manager = member(organization, Role.MANAGER)
    driver = member(organization, Role.DRIVER)

    answer = client_for(owner).get(reverse("report-partners")).json()

    assert (answer["totals"]["contributed"], answer["totals"]["share"], answer["totals"]["paid"]) == (
        "1000.000",
        "286.000",
        "286.000",
    )
    assert client_for(accountant).get(reverse("report-partners")).status_code == 200
    assert client_for(manager).get(reverse("report-partners")).status_code == 403
    assert client_for(driver).get(reverse("report-partners")).status_code == 403
    excel = client_for(owner).get(reverse("report-partners-export"))
    sheet = load_workbook(BytesIO(excel.content)).active
    assert sheet is not None
    assert sheet.cell(row=sheet.max_row, column=1).value == "Total"
